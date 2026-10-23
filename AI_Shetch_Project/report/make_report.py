from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT=Path(__file__).with_name('project_report.docx')
d=Document(); sec=d.sections[0]
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
styles=d.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9)
styles['Heading 1'].font.name='Aptos Display'; styles['Heading 1'].font.size=Pt(15); styles['Heading 1'].font.color.rgb=RGBColor(39,83,180)
styles['Heading 2'].font.size=Pt(11)

p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('PROJET ANI-2083'); r.bold=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor(39,83,180)
d.add_paragraph('\n')
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('AI Sketch'); r.bold=True; r.font.size=Pt(31); r.font.color.rgb=RGBColor(27,45,76)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run('Intelligent Drawing and Object Recognition Tool'); r.italic=True; r.font.size=Pt(15)
d.add_paragraph('\n')
for label,val in [('Cours','ANI-2083'),('Enseignant','Dr. Aurelle TCHAGNA'),('Période','13–19 octobre 2026'),('Auteurs','[À compléter : noms des membres du groupe]')]:
 p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run(label+': ').bold=True; p.add_run(val)
d.add_paragraph('\n\n')
p=d.add_paragraph('Rapport de projet — version de travail'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p=d.add_paragraph('Les résultats expérimentaux, captures d’écran et contributions doivent être complétés après exécution sur l’environnement du groupe.'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
d.add_page_break()

d.add_heading('1. Problème et objectifs',1)
d.add_paragraph("AI Sketch explore la classification d’images à partir de dessins simples. L’utilisateur trace une forme sur un canevas, puis l’application peut proposer une classe. Le projet vise à relier une interface de dessin, une préparation d’image cohérente et un classifieur convolutif reproductible.")
d.add_paragraph("Les cinq catégories retenues sont house (maison), tree (arbre), car (voiture), cat (chat) et sun (soleil). L’application permet également d’ajuster le pinceau, d’effacer et d’enregistrer le dessin en PNG.")
d.add_heading('2. Technologies et rôles',1)
d.add_paragraph("Python structure l’application et les scripts. PyQt5 fournit la fenêtre et le canevas interactif. Pillow et NumPy servent au traitement des images et aux tenseurs. TensorFlow/Keras définit et entraîne le CNN; scikit-learn réalise les partitions stratifiées et calcule les métriques; ndjson lit le format de données téléchargé.")
d.add_paragraph("L’environnement examiné est Windows 11, Python 3.14.4 64 bits. PyQt5, TensorFlow et Pillow étaient absents au début. TensorFlow nécessitant une version Python prise en charge, l’installation recommandée pour l’entraînement est Python 3.11 ou 3.12 dans un environnement virtuel.")
d.add_heading('3. Données et préparation',1)
d.add_paragraph("La source prévue est le jeu officiel Google Quick, Draw!, fichiers simplifiés NDJSON. train_model.py télécharge les catégories nécessaires et limite les exemples par classe. Chaque séquence de traits est rasterisée sur fond blanc, puis convertie en niveaux de gris, inversée, recadrée autour du trait, redimensionnée avec marge au format 64 × 64, centrée et normalisée entre 0 et 1. La même fonction preprocess est appelée par le pipeline de prédiction.")
d.add_paragraph("Les données téléchargées et le nombre effectivement lus doivent être consignés après exécution. La partition stratifiée est de 70 % entraînement, 15 % validation et 15 % test, avec graine fixe 42.")
d.add_heading('4. Architecture et fonctionnement',1)
d.add_paragraph("Le CNN Keras comporte deux blocs convolutionnels (24 puis 48 filtres, convolution 3 × 3, ReLU et max-pooling), une couche dense de 64 neurones, un dropout de 25 % et une sortie softmax à cinq classes. L’optimiseur Adam minimise l’entropie croisée catégorielle sparse; l’arrêt anticipé surveille la validation. Le modèle sauvegardé est ai_sketch_model.keras, avec l’ordre des classes conservé dans labels.json.")
d.add_paragraph("Flux applicatif : dessin souris → vérification de canevas non vide → prétraitement commun → scores du modèle chargé → classe de score maximal et score en pourcentage. Si le modèle manque, l’interface explique que la reconnaissance n’est pas disponible. Le score softmax n’est pas une garantie d’exactitude.")
d.add_heading('5. Tests et résultats mesurés',1)
d.add_paragraph("Vérifications réalisées sous Windows 11 et Python 3.14.4 : PyQt5 5.15.11, NumPy 2.5.4 et Pillow 13.0.0 installés; lancement de python app.py sans erreur. Un test Qt sans affichage a confirmé canevas vide au départ, détection d’un trait, enregistrement PNG réussi et canevas vide après effacement. Aucun entraînement ni résultat de précision n’a encore été mesuré. Aucune capture d’écran n’est incluse. Après installation de TensorFlow compatible, exécuter python train_model.py --per-class 1000; le script doit produire training_metrics.json avec les effectifs et la précision du jeu de test. Compléter à partir de ce fichier, sans estimation.")
table=d.add_table(rows=1, cols=2); table.style='Light Shading Accent 1'; table.rows[0].cells[0].text='Mesure'; table.rows[0].cells[1].text='Valeur observée à compléter'
for a,b in [('Exemples entraînement / validation / test','[À compléter]'),('Exactitude sur test','[À compléter]'),('Dessin vide refusé','[À compléter après essai]'),('Effacement et PNG','[À compléter après essai]')]:
 c=table.add_row().cells; c[0].text=a; c[1].text=b
d.add_heading('6. Limites, difficultés et améliorations',1)
d.add_paragraph("Le classifieur ne connaît que cinq classes et choisit nécessairement l’une d’elles, même pour un objet inconnu. Le score softmax n’est pas calibré. Les doodles Quick, Draw! diffèrent parfois des dessins libres; un petit sous-ensemble peut généraliser faiblement. La compatibilité TensorFlow dépend de Python et de la plateforme. Il faut vérifier les erreurs par classe et tester des dessins indépendants du jeu téléchargé.")
d.add_paragraph("Améliorations : ajouter des dessins originaux avec consentement, équilibrer les échantillons, examiner la matrice de confusion, calibrer un seuil de rejet sur validation, et documenter plusieurs prédictions réussies et échouées avec de vraies captures.")
d.add_heading('7. Conclusion et répartition du travail',1)
d.add_paragraph("Le prototype rassemble les composants essentiels d’une démonstration de classification de doodles. La remise finale doit être accompagnée du modèle réellement entraîné, des métriques issues du jeu de test et des preuves visuelles obtenues sur l’ordinateur de l’équipe.")
table=d.add_table(rows=1, cols=2); table.style='Light Shading Accent 1'; table.rows[0].cells[0].text='Rôle'; table.rows[0].cells[1].text='Membre / contribution à compléter'
for a in ['Interface et canevas','Données et préparation','Modèle et évaluation','Intégration, tests et rapport']:
 c=table.add_row().cells; c[0].text=a; c[1].text='[Nom et travail réellement effectué]'
d.save(OUT)
print(OUT)
